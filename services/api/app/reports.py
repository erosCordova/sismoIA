from app.time_utils import utc_iso, as_utc
import csv
import io
from datetime import datetime, timedelta, timezone
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import case, func, select
from sqlalchemy.orm import Session
from app.access import requerir_permiso
from app.database import engine
from app.models import DataSource, SeismicEvent
from typing import Literal

router = APIRouter(prefix="/api/v1/reports", tags=["Reportes"])


@router.get('/modules.csv')
def exportar_modulo(tipo: Literal['calidad', 'modelos', 'alertas'], dias: int = Query(30, ge=1, le=365),
                     peru: bool = True, fuente: Literal['IGP', 'USGS'] | None = None,
                     usuario=Depends(requerir_permiso('reports:create'))):
    from app.governance_models import ExperimentRun, DatasetVersion
    from app.alert_models import SeismicAlert, AlertRule
    cutoff = datetime.now(timezone.utc)
    conditions = filtros(dias, peru, fuente, cutoff)
    with Session(engine) as db:
        if tipo == 'calidad':
            rows = db.execute(select(DataSource.name, func.count(), func.count(SeismicEvent.depth_km),
                                     func.min(SeismicEvent.occurred_at), func.max(SeismicEvent.occurred_at))
                              .join(DataSource, DataSource.id == SeismicEvent.source_id)
                              .where(*conditions).group_by(DataSource.name)).all()
            columns = ['fuente', 'registros', 'con_profundidad', 'sin_profundidad', 'primer_evento_utc', 'ultimo_evento_utc']
            values = [[s, n, depth, n-depth, utc_iso(first), utc_iso(last)] for s,n,depth,first,last in rows]
        elif tipo == 'modelos':
            query = select(ExperimentRun.id, DatasetVersion.id, DatasetVersion.signature, ExperimentRun.algorithm,
                           ExperimentRun.results, ExperimentRun.status, ExperimentRun.created_at, ExperimentRun.artifact_sha256)
            query = query.join(DatasetVersion, DatasetVersion.id == ExperimentRun.dataset_id)
            criteria = [ExperimentRun.created_at.between(cutoff-timedelta(days=dias), cutoff),
                        func.coalesce(ExperimentRun.parameters['peru'].as_boolean(), DatasetVersion.manifest['peru'].as_boolean()) == peru]
            if fuente:
                criteria.append(func.coalesce(ExperimentRun.parameters['fuente'].as_string(), DatasetVersion.manifest['fuente'].as_string()) == fuente)
            rows = db.execute(query.where(*criteria).order_by(ExperimentRun.created_at.desc()).limit(10001)).all()
            columns = ['experimento_id', 'dataset_id', 'dataset_sha256', 'algoritmo', 'modelo', 'estado', 'fecha_utc', 'artefacto_sha256']
            values = [[str(eid), str(did), signature, algorithm, result.get('modelo_seleccionado') or result.get('modelo'),
                       status, utc_iso(created), artifact_hash] for eid,did,signature,algorithm,result,status,created,artifact_hash in rows]
        else:
            rows = db.execute(select(SeismicAlert, AlertRule.name, DataSource.name, SeismicEvent.magnitude)
                              .join(SeismicEvent, SeismicEvent.id == SeismicAlert.event_id)
                              .join(DataSource, DataSource.id == SeismicEvent.source_id)
                              .join(AlertRule, AlertRule.id == SeismicAlert.rule_id)
                              .where(*conditions, SeismicAlert.created_at <= cutoff)
                              .order_by(SeismicAlert.created_at.desc()).limit(10001)).all()
            columns = ['alerta_id', 'evento_id', 'regla', 'fuente', 'magnitud', 'severidad', 'estado', 'creada_utc', 'revisada_utc']
            values = [[str(a.id), str(a.event_id), rule, source, mag, a.severity, a.status, utc_iso(a.created_at), utc_iso(a.reviewed_at)] for a,rule,source,mag in rows]
    limited = len(values) > 10000
    stream = io.StringIO()
    writer = csv.writer(stream)
    context_columns = ['tipo_reporte', 'fecha_corte_utc', 'dias', 'peru', 'fuente_filtro', 'solicitante_id', 'limitado', 'filas_reporte']
    context = [tipo, utc_iso(cutoff), dias, peru, fuente or 'Todas', usuario.get('id'), limited, min(len(values),10000)]
    writer.writerow(context_columns + columns)
    for row in values[:10000] or [[None]*len(columns)]:
        writer.writerow([texto_seguro(value) if isinstance(value,str) else value for value in context+row])
    registrar_descarga(usuario, f'{tipo}.csv', dias, peru, fuente)
    return Response('\ufeff'+stream.getvalue(), media_type='text/csv; charset=utf-8', headers={
        'Content-Disposition':f'attachment; filename=sismoia_{tipo}.csv', 'Cache-Control':'no-store',
        'X-Export-Truncated':str(limited).lower()})
LIMITE_EXPORTACION = 10000

def filtros(dias, peru, fuente, corte=None):
    ahora = corte or datetime.now(timezone.utc)
    condiciones = [
        SeismicEvent.occurred_at >= ahora - timedelta(days=dias),
        SeismicEvent.occurred_at <= ahora
    ]
    if peru:
        condiciones += [
            SeismicEvent.latitude.between(-19, 1),
            SeismicEvent.longitude.between(-82, -68)
        ]
    if fuente:
        condiciones.append(DataSource.name == fuente)
    return condiciones

@router.get("/summary")
def resumen(
    dias: int = Query(30, ge=1, le=365),
    peru: bool = True,
    fuente: str | None = Query(None, max_length=100),
    usuario=Depends(requerir_permiso("reports:read"))
):
    corte = datetime.now(timezone.utc)
    condiciones = filtros(dias, peru, fuente, corte)
    rangos = [
        "Menor a 2", "2 a 3", "3 a 4", "4 a 5",
        "5 a 6", "6 a 7", "7 o mas"
    ]
    magnitud = case(
        (SeismicEvent.magnitude < 2, rangos[0]),
        (SeismicEvent.magnitude < 3, rangos[1]),
        (SeismicEvent.magnitude < 4, rangos[2]),
        (SeismicEvent.magnitude < 5, rangos[3]),
        (SeismicEvent.magnitude < 6, rangos[4]),
        (SeismicEvent.magnitude < 7, rangos[5]),
        else_=rangos[6]
    )
    with Session(engine) as db:
        total, media, maxima, desde, hasta = db.execute(
            select(func.count(SeismicEvent.id),
                   func.avg(SeismicEvent.magnitude),
                   func.max(SeismicEvent.magnitude),
                   func.min(SeismicEvent.occurred_at),
                   func.max(SeismicEvent.occurred_at))
            .select_from(SeismicEvent).join(DataSource)
            .where(*condiciones)
        ).one()
        por_fuente = db.execute(
            select(DataSource.name, func.count(SeismicEvent.id))
            .join(SeismicEvent, SeismicEvent.source_id == DataSource.id)
            .where(*condiciones)
            .group_by(DataSource.name).order_by(DataSource.name)
        ).all()
        mes = func.to_char(func.timezone("UTC", SeismicEvent.occurred_at), "YYYY-MM")
        por_mes = db.execute(
            select(mes.label("mes"), func.count(SeismicEvent.id))
            .select_from(SeismicEvent).join(DataSource)
            .where(*condiciones).group_by(mes).order_by(mes)
        ).all()
        por_magnitud = dict(db.execute(
            select(magnitud, func.count(SeismicEvent.id))
            .select_from(SeismicEvent).join(DataSource)
            .where(*condiciones).group_by(magnitud)
        ).all())
    return {
        "dias": dias, "peru": peru, "fuente": fuente,
        "fecha_corte_utc": corte.isoformat(),
        "solicitante": usuario.get("nombre") or usuario.get("id"),
        "solicitante_id": usuario.get("id"),
        "total": int(total),
        "magnitud_promedio": round(float(media), 2) if media is not None else None,
        "magnitud_maxima": float(maxima) if maxima is not None else None,
        "primera_fecha_utc": utc_iso(desde),
        "ultima_fecha_utc": utc_iso(hasta),
        "por_fuente": [{"nombre": n, "cantidad": int(c)} for n, c in por_fuente],
        "por_mes": [{"etiqueta": m, "cantidad": int(c)} for m, c in por_mes],
        "por_magnitud": [{"etiqueta": r, "cantidad": int(por_magnitud.get(r, 0))} for r in rangos],
        "advertencia": (
            "Estadisticas descriptivas del catalogo disponible. "
            "No predicen terremotos. USGS e IGP pueden contener eventos coincidentes. "
            "El filtro Peru utiliza un rectangulo geografico aproximado.")
    }

def texto_seguro(valor):
    texto = str(valor or "")
    if texto and texto[0] in ("=", "+", "-", "@", "\t", "\r"):
        return chr(39) + texto
    return texto

@router.get("/export.csv")
def exportar(
    dias: int = Query(30, ge=1, le=365),
    peru: bool = True,
    fuente: str | None = Query(None, max_length=100),
    usuario=Depends(requerir_permiso("reports:create"))
):
    corte = datetime.now(timezone.utc)
    condiciones = filtros(dias, peru, fuente, corte)
    with Session(engine) as db:
        total_disponible = int(db.scalar(select(func.count(SeismicEvent.id)).join(DataSource, DataSource.id == SeismicEvent.source_id).where(*condiciones)) or 0)
    registrar_descarga(usuario, "csv", dias, peru, fuente)
    consulta = (
        select(SeismicEvent.external_id, SeismicEvent.occurred_at,
               DataSource.name, SeismicEvent.magnitude,
               SeismicEvent.depth_km, SeismicEvent.latitude,
               SeismicEvent.longitude, SeismicEvent.place,
               SeismicEvent.source_url)
        .join(DataSource, DataSource.id == SeismicEvent.source_id)
        .where(*condiciones)
        .order_by(SeismicEvent.occurred_at.desc())
        .limit(LIMITE_EXPORTACION)
    )
    def generar():
        salida = io.StringIO()
        escritor = csv.writer(salida)
        escritor.writerow([
            "id_externo", "fecha_utc", "fuente", "magnitud",
            "profundidad_km", "latitud", "longitud", "lugar", "url_fuente",
            "fecha_corte_utc", "filtro_dias", "filtro_peru", "filtro_fuente", "solicitante_id", "total_disponible", "exportacion_limitada"
        ])
        yield "\ufeff" + salida.getvalue()
        with Session(engine) as db:
            for fila in db.execute(consulta):
                salida.seek(0)
                salida.truncate(0)
                id_externo, fecha, origen, mag, profundidad, lat, lon, lugar, url = fila
                escritor.writerow([
                    texto_seguro(id_externo), utc_iso(fecha),
                    texto_seguro(origen), mag, profundidad, lat, lon,
                    texto_seguro(lugar), texto_seguro(url),
                    corte.isoformat(), dias, peru, texto_seguro(fuente or "Todas"),
                    texto_seguro(usuario.get("id")), total_disponible, total_disponible > LIMITE_EXPORTACION
                ])
                yield salida.getvalue()
    return StreamingResponse(
        generar(), media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": "attachment; filename=sismoia_reporte.csv",
            "X-Export-Total": str(total_disponible),
            "X-Export-Truncated": str(total_disponible > LIMITE_EXPORTACION).lower(),
            "Cache-Control": "no-store"
        }
    )


# SISMOIA_REPORT_FORMATS_V2

from uuid import UUID
from fastapi.responses import Response
from app.models import AuditLog


def registrar_descarga(usuario, formato, dias, peru, fuente):
    try:
        identificador = UUID(str(usuario["id"]))
    except (KeyError, TypeError, ValueError):
        from fastapi import HTTPException
        raise HTTPException(
            403,
            "Se requiere identidad autenticada para exportar"
        )

    with Session(engine) as db:
        db.add(AuditLog(
            actor_id=identificador,
            action="reports.export",
            entity_type="report",
            entity_id=formato,
            after_json={
                "formato": formato,
                "dias": dias,
                "peru": peru,
                "fuente": fuente,
                "limite_csv": LIMITE_EXPORTACION,
            }
        ))
        db.commit()


def datos_resumen(dias, peru, fuente, usuario):
    return resumen(
        dias=dias,
        peru=peru,
        fuente=fuente,
        usuario=usuario
    )


@router.get("/export.xlsx")
def exportar_xlsx(
    dias: int = Query(30, ge=1, le=365),
    peru: bool = True,
    fuente: str | None = Query(
        None, pattern="^(USGS|IGP)$"
    ),
    usuario=Depends(requerir_permiso("reports:create")),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    datos = datos_resumen(
        dias, peru, fuente, usuario
    )

    libro = Workbook()
    principal = libro.active
    principal.title = "Resumen"

    def texto_excel(v):
        if isinstance(v, str) and v.startswith(
            ("=", "+", "-", "@")
        ):
            return "'" + v
        return v

    def hoja(nombre, cabeceras, filas):
        if nombre == "Resumen":
            ws = principal
        else:
            ws = libro.create_sheet(nombre)

        ws.append(cabeceras)

        for fila in filas:
            ws.append([
                texto_excel(v) for v in fila
            ])

        for celda in ws[1]:
            celda.font = Font(
                bold=True, color="FFFFFF"
            )
            celda.fill = PatternFill(
                "solid", fgColor="1664A6"
            )
            celda.alignment = Alignment(
                vertical="center"
            )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for columna in ws.columns:
            valores = [
                len(str(c.value or ""))
                for c in columna
            ]
            ancho = min(
                46, max(14, max(valores) + 3)
            )
            ws.column_dimensions[
                columna[0].column_letter
            ].width = ancho

        return ws

    hoja("Resumen", ["Indicador", "Valor"], [
        ("Sistema", "SISMOIA Enterprise"),
        ("Solicitante", datos["solicitante"]),
        ("Fecha de corte UTC", datos["fecha_corte_utc"]),
        ("Contenido", "Resumen agregado; CSV contiene eventos individuales"),
        ("Zona", "Perú" if peru else "Mundo"),
        ("Fuente", fuente or "Todas"),
        ("Periodo (días)", dias),
        ("Registros", datos["total"]),
        ("Magnitud media", datos["magnitud_promedio"]),
        ("Magnitud máxima", datos["magnitud_maxima"]),
        ("Primer evento UTC", datos["primera_fecha_utc"]),
        ("Último evento UTC", datos["ultima_fecha_utc"]),
        (
            "Generado UTC",
            datetime.now(timezone.utc).isoformat()
        ),
        ("Advertencia", datos["advertencia"]),
    ])

    for titulo, contenido, clave in [
        ("Magnitudes", datos["por_magnitud"], "etiqueta"),
        ("Meses", datos["por_mes"], "etiqueta"),
        ("Fuentes", datos["por_fuente"], "nombre"),
    ]:
        hoja(
            titulo,
            [titulo[:-1] if titulo.endswith("s") else titulo, "Registros"],
            [
                (fila[clave], fila["cantidad"])
                for fila in contenido
            ]
        )

    salida = io.BytesIO()
    libro.save(salida)

    registrar_descarga(
        usuario, "xlsx", dias, peru, fuente
    )

    return Response(
        content=salida.getvalue(),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                "attachment; filename=sismoia_reporte.xlsx",
            "Cache-Control": "no-store",
        }
    )


@router.get("/export.pdf")
def exportar_pdf(
    dias: int = Query(30, ge=1, le=365),
    peru: bool = True,
    fuente: str | None = Query(
        None, pattern="^(USGS|IGP)$"
    ),
    usuario=Depends(requerir_permiso("reports:create")),
):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4

    datos = datos_resumen(
        dias, peru, fuente, usuario
    )

    salida = io.BytesIO()
    pdf = canvas.Canvas(
        salida, pagesize=A4
    )

    ancho, alto = A4

    pdf.setFillColorRGB(0.08, 0.30, 0.52)
    pdf.rect(
        0, alto - 110,
        ancho, 110,
        stroke=0, fill=1
    )

    pdf.setFillColorRGB(1, 1, 1)
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(
        42, alto - 54,
        "SISMOIA Enterprise"
    )

    pdf.setFont("Helvetica", 11)
    pdf.drawString(
        42, alto - 78,
        "Reporte de actividad sismica"
    )

    pdf.setFillColorRGB(0.15, 0.29, 0.42)

    y = alto - 142

    def linea(etiqueta, valor):
        nonlocal_y = None
        return f"{etiqueta}: {valor}"

    lineas = [
        ("Solicitante", str(datos["solicitante"])[:65]),
        ("Corte UTC", datos["fecha_corte_utc"]),
        ("Zona", "Peru" if peru else "Mundo"),
        ("Fuente", fuente or "Todas"),
        ("Periodo", f"{dias} dias"),
        ("Registros", datos["total"]),
        (
            "Magnitud media",
            datos["magnitud_promedio"]
        ),
        (
            "Magnitud maxima",
            datos["magnitud_maxima"]
        ),
        (
            "Fecha de generacion UTC",
            datetime.now(timezone.utc).strftime(
                "%Y-%m-%d %H:%M"
            )
        ),
    ]

    pdf.setFont("Helvetica", 11)

    for etiqueta, valor in lineas:
        pdf.drawString(
            42, y,
            linea(etiqueta, valor)
        )
        y -= 23

    y -= 12
    pdf.setFont("Helvetica-Bold", 13)
    pdf.drawString(
        42, y,
        "Distribucion por magnitud"
    )
    y -= 26

    distribucion = datos["por_magnitud"]

    maximo = max(
        1,
        *(f["cantidad"] for f in distribucion)
    )

    for fila in distribucion:
        cantidad = fila["cantidad"]
        etiqueta = fila["etiqueta"]

        pdf.setFont("Helvetica", 9)
        pdf.drawString(
            42, y, str(etiqueta)
        )

        largo = 270 * cantidad / maximo

        pdf.setFillColorRGB(
            0.09, 0.59, 0.74
        )
        pdf.rect(
            155, y - 3,
            largo, 10,
            stroke=0, fill=1
        )

        pdf.setFillColorRGB(
            0.15, 0.29, 0.42
        )
        pdf.drawString(
            439, y,
            str(cantidad)
        )

        y -= 26

    y -= 12

    pdf.setFont("Helvetica-Bold", 13)
    pdf.drawString(
        42, y,
        "Registros por fuente"
    )

    y -= 24
    pdf.setFont("Helvetica", 10)

    for fila in datos["por_fuente"]:
        pdf.drawString(
            42, y,
            f'{fila["nombre"]}: {fila["cantidad"]}'
        )
        y -= 20

    y = max(77, y - 22)

    pdf.setFont("Helvetica", 9)

    advertencias = [
        "Estadisticas descriptivas de registros almacenados.",
        "No representan predicciones de futuros terremotos.",
        "USGS e IGP pueden registrar un mismo evento.",
        "La delimitacion de Peru es geografica aproximada.",
    ]

    for aviso in advertencias:
        pdf.drawString(
            42, y, aviso
        )
        y -= 14

    pdf.setStrokeColorRGB(
        0.78, 0.85, 0.90
    )

    pdf.line(
        42, 53,
        ancho - 42, 53
    )

    pdf.setFont("Helvetica", 8)
    pdf.drawString(
        42, 39,
        "SISMOIA - Documento generado automaticamente"
    )

    pdf.showPage()
    pdf.save()

    registrar_descarga(
        usuario, "pdf", dias, peru, fuente
    )

    return Response(
        content=salida.getvalue(),
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                "attachment; filename=sismoia_reporte.pdf",
            "Cache-Control": "no-store",
        }
    )
