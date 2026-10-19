"""Integracion sobre PostgreSQL efimero explicitamente habilitado; rollback por caso."""
import os
import csv
import io
import subprocess
import sys
import unittest
import tempfile
from contextlib import ExitStack
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import Mock, patch
from uuid import UUID

from sqlalchemy import select, text
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session
from fastapi.testclient import TestClient

from app import alerts, deliveries, governance, seismic, ingestion, preparation, reports, science, ml, management, auth, zone_analysis, seismic_forecast_api, email_notifications
from app.auth import usuario_actual
from app.database import engine, settings
from app.main import app
from app.models import User, Role, UserRole, DataSource, SeismicEvent
from app.ingestion import IngestionRun

TEST_URL = os.getenv("SISMOIA_TEST_POSTGRES_URL", "")


@unittest.skipUnless(TEST_URL, "Requiere PostgreSQL efimero y SISMOIA_TEST_POSTGRES_URL explicito")
class PlanPostgresTests(unittest.TestCase):
    @classmethod
    def tearDownClass(cls):
        engine.dispose()

    @classmethod
    def setUpClass(cls):
        url = make_url(TEST_URL)
        if url.host not in ("127.0.0.1", "localhost") or not (url.database or "").startswith("sismoia_test"):
            raise RuntimeError("Las pruebas solo admiten una base local con prefijo sismoia_test")
        if make_url(settings.database_url) != url:
            raise RuntimeError("DATABASE_URL debe coincidir con la base aislada de pruebas")
        root = Path(__file__).resolve().parents[3]
        subprocess.run([sys.executable, "-m", "alembic", "upgrade", "head"], cwd=root, check=True, capture_output=True)

    def setUp(self):
        self.conn = engine.connect()
        self.transaction = self.conn.begin()
        self.stack = ExitStack()

        def session(*args, **kwargs):
            return Session(bind=self.conn, join_transaction_mode="create_savepoint")

        self.session = session
        for module in (alerts, deliveries, governance, seismic, ingestion, preparation, reports, science, management, auth, zone_analysis, seismic_forecast_api, email_notifications):
            self.stack.enter_context(patch.object(module, "Session", side_effect=session))
        with session() as db:
            role = Role(name="Administrador")
            user = User(full_name="Prueba aislada", email="test@example.invalid", password_hash="!sin-login!", status="active")
            db.add_all([role, user])
            db.flush()
            db.add(UserRole(user_id=user.id, role_id=role.id))
            self.user_id = user.id
            db.commit()
        app.dependency_overrides[usuario_actual] = lambda: {"id": str(self.user_id), "roles": ["Administrador"]}
        self.client = self.stack.enter_context(TestClient(app))

    def tearDown(self):
        app.dependency_overrides.clear()
        self.stack.close()
        if self.transaction.is_active:
            self.transaction.rollback()
        self.conn.close()

    def test_agrupamientos_y_regresion_desde_postgres_sin_csv(self):
        now = datetime.now(timezone.utc)
        with self.session() as db:
            source = DataSource(name="IGP", organization="Prueba", base_url="https://example.invalid", source_type="test")
            db.add(source)
            db.flush()
            for i in range(80):
                db.add(SeismicEvent(source_id=source.id, external_id=f"ml-{i}", occurred_at=now-timedelta(days=81-i),
                                   latitude=-18+i%10, longitude=-80+i%7, magnitude=3+i%4/2, depth_km=20+i%40))
            db.commit()
        with tempfile.TemporaryDirectory() as folder, patch.object(ml, 'BASE', Path(folder)):
            self.assertFalse(self.client.get('/api/v1/ml/overview').json()['disponible'])
            response = self.client.post('/api/v1/ml/train-clustering?peru=true')
            self.assertEqual(response.status_code, 200, response.text)
            self.assertTrue(self.client.get('/api/v1/ml/overview').json()['disponible'])
            self.assertGreater(len(self.client.get('/api/v1/ml/clusters').json()['puntos']), 0)
            response = self.client.get('/api/v1/ml/regression-validation?peru=true')
            self.assertEqual(response.status_code, 200, response.text)
            self.assertTrue(response.json()['peru'])
            response = self.client.post('/api/v1/science/ml/compare?peru=true')
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(len(response.json()['regresion']), 3)

    def test_zonas_dataset_revision_publicacion_y_reportes_coherentes(self):
        now = datetime.now(timezone.utc)
        with self.session() as db:
            source = DataSource(name='IGP',organization='Prueba',base_url='https://example.invalid',source_type='test')
            db.add(source); db.flush()
            db.add(SeismicEvent(source_id=source.id,external_id='outside-peru',occurred_at=now-timedelta(days=2),
                               latitude=35,longitude=139,magnitude=6,depth_km=20))
            for i in range(80):
                db.add(SeismicEvent(source_id=source.id,external_id=f'zone-{i}',occurred_at=now-timedelta(days=81-i),
                                   latitude=-12,longitude=-77,magnitude=3+i%4/2,depth_km=20))
            db.commit()
        live = self.client.get('/api/v1/forecasts/zones?fuente=IGP&peru=true&dias=365')
        self.assertEqual(live.status_code,200,live.text)
        self.assertEqual(live.json()['registros_historicos'],80)
        self.assertEqual(live.json()['registros_umbral'],40)
        self.assertFalse(live.json()['metadatos']['consulta_limitada'])
        world=self.client.get('/api/v1/forecasts/zones?fuente=IGP&peru=false&dias=365').json()
        self.assertEqual(world['registros_historicos'],81)
        self.assertEqual(len(world['zonas']),24)
        response = self.client.post('/api/v1/governance/datasets',json={'nombre':'Zonas de prueba','fuente':'IGP','limite':100})
        self.assertEqual(response.status_code,201,response.text)
        did = response.json()['id']
        self.assertEqual(self.client.post('/api/v1/forecasts/runs',json={'dataset_id':did,'fuente':'IGP','peru':False}).status_code,422)
        response = self.client.post('/api/v1/forecasts/runs',json={'dataset_id':did,'fuente':'IGP','peru':True})
        self.assertEqual(response.status_code,201,response.text)
        eid = response.json()['id']
        self.assertEqual(self.client.get(f'/api/v1/forecasts/published/{eid}').status_code,409)
        decision = {'decision':'approved','motivo':'Revision experimental en prueba aislada'}
        self.assertEqual(self.client.post(f'/api/v1/governance/experiments/{eid}/review',json=decision).status_code,409)
        self.assertEqual(self.client.post(f'/api/v1/governance/datasets/{did}/review',json=decision).status_code,200)
        self.assertEqual(self.client.post(f'/api/v1/governance/experiments/{eid}/review',json=decision).status_code,200)
        published = self.client.get(f'/api/v1/forecasts/published/{eid}')
        self.assertEqual(published.status_code,200,published.text)
        self.assertTrue(published.json()['publicado'])
        self.assertEqual(published.json()['dataset_id'],did)
        self.assertIn(eid,[r['id'] for r in self.client.get('/api/v1/forecasts/published').json()['versiones']])
        self.assertNotIn(eid,[r['id'] for r in self.client.get('/api/v1/governance/experiments').json()['experimentos']])
        app.dependency_overrides[usuario_actual]=lambda:{'id':str(self.user_id),'roles':['Solo lectura']}
        self.assertEqual(self.client.get(f'/api/v1/forecasts/published/{eid}').status_code,200)
        self.assertEqual(self.client.post('/api/v1/forecasts/runs',json={'dataset_id':did,'fuente':'IGP','peru':True}).status_code,403)
        app.dependency_overrides[usuario_actual]=lambda:{'id':str(self.user_id),'roles':['Administrador']}
        for kind in ('calidad','modelos','alertas'):
            response = self.client.get(f'/api/v1/reports/modules.csv?tipo={kind}&dias=365&peru=true')
            self.assertEqual(response.status_code,200,response.text)
            self.assertIn('solicitante_id',response.text)
            if kind == 'modelos': self.assertIn(eid,response.text)
        self.assertEqual(self.client.get('/api/v1/management/audit/export').status_code,200)
        with self.session() as db:
            from app.governance_models import ExperimentRun
            db.get(ExperimentRun,UUID(eid)).artifact=b'alterado'
            db.commit()
        self.assertEqual(self.client.get(f'/api/v1/forecasts/published/{eid}').status_code,409)

    def test_estado_cuentas_y_cierre_real_de_sesion(self):
        from app.passwords import hash_password
        app.dependency_overrides.clear()
        with self.session() as db:
            db.get(User,self.user_id).password_hash=hash_password('PruebaAislada123')
            role=Role(name='Analista'); other=User(full_name='Analista prueba',email='analyst@example.invalid',password_hash=hash_password('PruebaAislada123'),status='active')
            db.add_all([role,other]); db.flush()
            other_id=other.id
            db.add(UserRole(user_id=other.id,role_id=role.id)); db.commit()
        admin=self.client.post('/api/v1/auth/login',json={'email':'test@example.invalid','password':'PruebaAislada123'}).json()['access_token']
        token=self.client.post('/api/v1/auth/login',json={'email':'analyst@example.invalid','password':'PruebaAislada123'}).json()['access_token']
        ah={'Authorization':f'Bearer {admin}'}; uh={'Authorization':f'Bearer {token}'}
        self.assertEqual(self.client.get('/api/v1/auth/me',headers=uh).status_code,200)
        self.assertEqual(self.client.patch(f'/api/v1/management/users/{other_id}/status',json={'estado':'inactive'},headers=ah).status_code,200)
        self.assertEqual(self.client.get('/api/v1/auth/me',headers=uh).status_code,401)
        self.assertEqual(self.client.patch(f'/api/v1/management/users/{other_id}/status',json={'estado':'active'},headers=ah).status_code,200)
        self.assertEqual(self.client.post('/api/v1/auth/logout',headers=uh).status_code,204)
        self.assertEqual(self.client.get('/api/v1/auth/me',headers=uh).status_code,401)
        self.assertEqual(self.client.get('/api/v1/auth/me',headers=ah).status_code,200)
        self.assertEqual(self.client.patch(f'/api/v1/management/users/{self.user_id}/status',json={'estado':'inactive'},headers=ah).status_code,409)

    def test_correo_reintentos_controlados_sin_envio_externo(self):
        from app.alert_models import AlertRule, SeismicAlert, AlertAction
        with self.session() as db:
            source=DataSource(name='IGP',organization='Prueba',base_url='https://example.invalid',source_type='test')
            db.add(source); db.flush()
            event=SeismicEvent(source_id=source.id,external_id='smtp-test',occurred_at=datetime.now(timezone.utc),latitude=-12,longitude=-77,magnitude=5,depth_km=20)
            rule=AlertRule(name='Correo de prueba',min_magnitude=4)
            db.add_all([event,rule]); db.flush()
            alert=SeismicAlert(event_id=event.id,rule_id=rule.id,severity='warning',status='approved')
            db.add(alert); db.flush(); aid=alert.id; db.commit()
        payload={'alert_id':str(aid),'recipient_user_id':str(self.user_id)}
        with patch.object(email_notifications,'configuracion',return_value={'habilitado':True,'configurado':True}), patch.object(email_notifications,'enviar_smtp',side_effect=RuntimeError('Simulacion sin red')) as smtp:
            self.assertEqual(self.client.post('/api/v1/email-notifications/send',json=payload).status_code,502)
            self.assertEqual(self.client.post('/api/v1/email-notifications/send',json=payload).status_code,429)
            self.assertEqual(smtp.call_count,1)
        with self.session() as db:
            action=db.scalar(select(AlertAction).where(AlertAction.alert_id==aid,AlertAction.action=='email.failed'))
            action.created_at=datetime.now(timezone.utc)-timedelta(seconds=61); db.commit()
        with patch.object(email_notifications,'configuracion',return_value={'habilitado':True,'configurado':True}), patch.object(email_notifications,'enviar_smtp') as smtp:
            self.assertEqual(self.client.post('/api/v1/email-notifications/send',json=payload).status_code,200)
            self.assertEqual(self.client.post('/api/v1/email-notifications/send',json=payload).status_code,409)
            self.assertEqual(smtp.call_count,1)

    def test_migracion_alertas_entregas_y_ml(self):
        now = datetime.now(timezone.utc)
        with self.session() as db:
            igp = DataSource(name="IGP", organization="Prueba", base_url="https://example.invalid", source_type="test")
            usgs = DataSource(name="USGS", organization="Prueba", base_url="https://example.invalid", source_type="test")
            db.add_all([igp, usgs])
            db.flush()
            for i in range(60):
                db.add(SeismicEvent(source_id=igp.id, external_id=f"historic-{i}", occurred_at=now-timedelta(days=61-i),
                                   latitude=-12, longitude=-77, magnitude=3 if i % 2 else 5, depth_km=20+i))
            for source in (igp, usgs):
                db.add(SeismicEvent(source_id=source.id, external_id="recent", occurred_at=now-timedelta(minutes=1),
                                   latitude=-12, longitude=-77, magnitude=5, depth_km=20))
            db.commit()
        response = self.client.post("/api/v1/alerts/rules", json={"nombre":"Solo IGP", "magnitud_minima":4, "fuente_permitida":"IGP"})
        self.assertEqual(response.status_code, 201, response.text)
        self.assertEqual(response.json()["fuente_permitida"], "IGP")
        response = self.client.post("/api/v1/alerts/evaluate")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["nuevas_alertas"], 1)
        self.assertEqual(self.client.post("/api/v1/alerts/evaluate").json()["nuevas_alertas"], 0)
        alert = self.client.get("/api/v1/alerts/inbox").json()["alertas"][0]
        self.assertEqual(alert["fuente"], "IGP")
        aid = alert["id"]
        self.assertEqual(self.client.post(f"/api/v1/alerts/{aid}/review", json={"decision":"approved"}).status_code, 200)
        response = self.client.post("/api/v1/deliveries/assign", json={"alerta_id":aid,"destinatario_id":str(self.user_id)})
        self.assertEqual(response.status_code, 201, response.text)
        did = response.json()["id"]
        self.assertEqual(self.client.post(f"/api/v1/deliveries/{did}/ack").status_code, 200)
        detail = self.client.get(f"/api/v1/alerts/inbox/{aid}").json()
        self.assertEqual(detail["entrega"]["estado"], "acknowledged")

        response = self.client.post("/api/v1/governance/datasets", json={"nombre":"IGP aislado", "fuente":"IGP", "limite":100})
        self.assertEqual(response.status_code, 201, response.text)
        dataset_id = response.json()["id"]
        decision = {"decision":"approved", "motivo":"Validacion aislada de integracion"}
        self.assertEqual(self.client.post(f"/api/v1/governance/datasets/{dataset_id}/review", json=decision).status_code, 200)
        response = self.client.post("/api/v1/governance/experiments", json={"dataset_id":dataset_id,"algoritmo":"clasificacion"})
        self.assertEqual(response.status_code, 201, response.text)
        eid = response.json()["id"]
        self.assertEqual(self.client.post(f"/api/v1/governance/experiments/{eid}/review", json=decision).status_code, 200)
        response = self.client.post(f"/api/v1/governance/experiments/{eid}/predict", json={"latitud":-12,"longitud":-77,"profundidad_km":25})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("probabilidad_clase_positiva", response.json())
        with self.session() as db:
            for table in ("sismo_dataset_versions", "sismo_ml_experiments", "sismo_alert_deliveries", "ingestion_runs"):
                self.assertTrue(db.scalar(text("SELECT relrowsecurity FROM pg_class WHERE oid=to_regclass(:name)"), {"name":"public."+table}))

    def test_ingesta_idempotente_actualiza_revision_y_registra_fallo(self):
        import httpx
        feature = {"id":"event-test", "properties":{"mag":4.2,"time":int(datetime.now(timezone.utc).timestamp()*1000), "status":"automatic"},
                   "geometry":{"coordinates":[-77,-12,25]}}
        reply = Mock()
        reply.json.side_effect = lambda: {"features":[feature]}
        client = Mock()
        client.__enter__ = Mock(return_value=client)
        client.__exit__ = Mock(return_value=False)
        client.get.return_value = reply
        identity = {"roles":["Administrador"]}
        with patch.object(seismic.httpx, "Client", return_value=client):
            first = seismic.sincronizar(identity)
            second = seismic.sincronizar(identity)
            feature["properties"]["mag"] = 4.8
            revised = seismic.sincronizar(identity)
        self.assertEqual(first["importados"], 1)
        self.assertEqual(second["duplicados"], 1)
        self.assertEqual(revised["actualizados"], 1)
        with self.session() as db:
            event = db.scalar(select(SeismicEvent).where(SeismicEvent.external_id=="event-test"))
            self.assertEqual(float(event.magnitude), 4.8)
            self.assertEqual(len(event.origin_hash), 64)
            self.assertEqual(len(db.scalars(select(IngestionRun)).all()), 3)
        client.get.side_effect = httpx.RequestError("error simulado sin datos sensibles")
        with patch.object(seismic.httpx, "Client", return_value=client), self.assertRaises(Exception):
            seismic.sincronizar(identity)
        with self.session() as db:
            run = db.scalar(select(IngestionRun).order_by(IngestionRun.started_at.desc()))
            self.assertEqual(run.status, "failed")
            self.assertIn("HTTPException", run.error_summary)

    def test_reportes_corte_solicitante_y_correlacion(self):
        from openpyxl import load_workbook
        from app.models import AuditLog
        with self.session() as db:
            source = DataSource(name="IGP", organization="Prueba", base_url="https://example.invalid", source_type="test")
            db.add(source)
            db.flush()
            db.add(SeismicEvent(source_id=source.id, external_id="report-test", occurred_at=datetime.now(timezone.utc),
                               latitude=-12, longitude=-77, magnitude=4.5, depth_km=25, place="Lima"))
            db.commit()
        response = self.client.get("/api/v1/reports/export.csv?fuente=IGP", headers={"X-Request-ID":"reporte-test-123"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["X-Export-Total"], "1")
        self.assertEqual(response.headers["X-Export-Truncated"], "false")
        rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(rows[0]["solicitante_id"], str(self.user_id))
        self.assertEqual(rows[0]["filtro_fuente"], "IGP")
        self.assertTrue(rows[0]["fecha_utc"].endswith("Z"))
        response = self.client.get("/api/v1/reports/export.xlsx?fuente=IGP")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        values = dict(workbook["Resumen"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(values["Solicitante"], str(self.user_id))
        self.assertIn("Fecha de corte UTC", values)
        workbook.close()
        response = self.client.get("/api/v1/reports/export.pdf?fuente=IGP")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF-"))
        if os.getenv("SISMOIA_TEST_PDF_OUTPUT"):
            Path(os.environ["SISMOIA_TEST_PDF_OUTPUT"]).write_bytes(response.content)
        with self.session() as db:
            audit = db.scalar(select(AuditLog).where(AuditLog.action=="reports.export", AuditLog.entity_id=="csv"))
            self.assertEqual(audit.correlation_id, "reporte-test-123")
